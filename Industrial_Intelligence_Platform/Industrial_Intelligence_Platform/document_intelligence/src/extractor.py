"""Document reading and pretrained question-answering metadata extraction."""
from pathlib import Path
from typing import Dict, Tuple

FIELDS = {
    "agreement_value": "What is the agreement value or monthly rent amount?",
    "agreement_start_date": "What is the start date of the rental agreement?",
    "agreement_end_date": "What is the end date of the rental agreement?",
    "renewal_notice_days": "How many days of notice are required for renewal or termination?",
    "party_one": "Who is the first party to the rental agreement, such as the landlord or owner?",
    "party_two": "Who is the second party to the rental agreement, such as the tenant or renter?",
}

def _read_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)

def _ocr_image(path: Path) -> str:
    from PIL import Image
    import pytesseract
    with Image.open(path) as im:
        return pytesseract.image_to_string(im)

def _read_pdf(path: Path) -> str:
    import fitz
    pages = []
    with fitz.open(str(path)) as pdf:
        for page in pdf:
            text = page.get_text("text").strip()
            if not text:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                from PIL import Image
                import pytesseract
                import io
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                text = pytesseract.image_to_string(image)
            pages.append(text)
    return "\n".join(pages)

def read_document(path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        return _ocr_image(path)
    raise ValueError(f"Unsupported file type: {suffix}")

class MetadataExtractor:
    """Extract six fields with a pretrained extractive QA transformer."""
    def __init__(self, model_name: str = "deepset/roberta-base-squad2", min_confidence: float = 0.08):
        from transformers import pipeline
        self.qa = pipeline("question-answering", model=model_name, tokenizer=model_name)
        self.min_confidence = min_confidence

    def extract_text(self, text: str, file_name: str = "") -> dict:
        if not text.strip():
            raise ValueError("No readable text was found in the document.")
        fields = {}
        for key, question in FIELDS.items():
            try:
                result = self.qa(question=question, context=text[:12000])
                answer = (result.get("answer") or "").strip()
                score = float(result.get("score", 0.0))
                if score < self.min_confidence:
                    answer = ""
                fields[key] = {"value": answer, "confidence": score}
            except Exception:
                fields[key] = {"value": "", "confidence": 0.0}
        return {"file_name": file_name, "fields": fields}

    def extract_file(self, path) -> dict:
        path = Path(path)
        text = read_document(path)
        return self.extract_text(text, path.name)
